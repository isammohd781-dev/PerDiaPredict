"""
apply_excel_accounts.py  -  يجعل الحسابات تُحفظ وتُفحص من ملف Excel (accounts.xlsx).

الاستخدام (ضعه بجانب diabetes_app.py ثم شغّله):
    python apply_excel_accounts.py

- يحفظ نسخة احتياطية: diabetes_app.py.bak4
- يستبدل قسم الحسابات (بين ACCOUNTS-DB-START و ACCOUNTS-DB-END) فقط.
- يضيف قسم "Registered accounts" داخل لوحة الإدارة (عرض + تحميل + استرجاع).
- يمكن تشغيله أكثر من مرة بدون مشاكل.
"""
import re
import shutil
import sys

APP_FILE = "diabetes_app.py"

NEW_BLOCK = r'''# ACCOUNTS-DB-START
# =============================================================================
# Accounts stored in an Excel file:  accounts.xlsx   (sheet "users")
#
# One row per registered user: ID, first / last name, email, birth date,
# password hash, created at, last login, failed attempts, locked until.
#
# * The password is NEVER written as readable text - only a salted
#   PBKDF2-SHA256 hash (a one-way code). At sign-in the typed password is
#   hashed the same way and compared with it.
# * The file is read into memory once and re-read only when it changes, so
#   the check at sign-in is instant.
# * Old accounts from perdiapredict.db (the previous storage) are imported
#   automatically the first time, if that file exists.
# =============================================================================
ACCOUNTS_FILE = "accounts.xlsx"
DB_FILE = "perdiapredict.db"          # old SQLite storage (only read for the one-time import)
PBKDF2_ITERATIONS = 260_000
MIN_PASSWORD_LEN = 8
MAX_FAILED_LOGINS = 5      # wrong passwords in a row ...
LOCK_MINUTES = 5           # ... lock that account for this many minutes
SUPPORT_WHATSAPP_NUMBER = "+256771715275"
SUPPORT_WHATSAPP_URL = "https://wa.me/256771715275"
EMAIL_REGEX = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")

ACCOUNT_FIELDS = [
    "id", "first_name", "last_name", "email", "birth_date",
    "password_hash", "created_at", "last_login", "failed_attempts", "locked_until",
]
ACCOUNT_HEADERS = [
    "ID", "First name", "Last name", "Email", "Birth date",
    "Password hash", "Created at", "Last login", "Failed attempts", "Locked until",
]
_FIELD_BY_HEADER = dict(zip(ACCOUNT_HEADERS, ACCOUNT_FIELDS))


@st.cache_resource
def _accounts_store():
    """Shared by every browser session: a lock + the in-memory copy of the file."""
    import threading

    return {"lock": threading.RLock(), "mtime": None, "rows": []}


def normalize_email(email: str) -> str:
    return (email or "").strip().lower()


def hash_password(password: str) -> str:
    salt = secrets.token_bytes(16)
    digest = hashlib.pbkdf2_hmac("sha256", password.encode("utf-8"), salt, PBKDF2_ITERATIONS)
    return f"pbkdf2_sha256${PBKDF2_ITERATIONS}${salt.hex()}${digest.hex()}"


def verify_password(password: str, stored: str) -> bool:
    try:
        _algo, iterations, salt_hex, hash_hex = stored.split("$")
        digest = hashlib.pbkdf2_hmac(
            "sha256", password.encode("utf-8"), bytes.fromhex(salt_hex), int(iterations)
        )
        return hmac.compare_digest(digest.hex(), hash_hex)
    except Exception:
        return False


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.isoformat(timespec="seconds")
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def _rows_from_sheet(ws) -> list:
    """Read the accounts out of a worksheet (columns are found by their header names)."""
    rows_iter = ws.iter_rows(values_only=True)
    header = next(rows_iter, None)
    if not header:
        return []

    index = {}
    for i, name in enumerate(header):
        field = _FIELD_BY_HEADER.get(_cell_text(name))
        if field:
            index[field] = i
    if "email" not in index or "password_hash" not in index:
        return []

    def get(raw, field):
        i = index.get(field)
        return _cell_text(raw[i]) if i is not None and i < len(raw) else ""

    def as_int(text):
        try:
            return int(float(text))
        except (TypeError, ValueError):
            return 0

    out = []
    for raw in rows_iter:
        email = normalize_email(get(raw, "email"))
        pw_hash = get(raw, "password_hash")
        if not email or not pw_hash:
            continue
        out.append({
            "id": as_int(get(raw, "id")),
            "first_name": get(raw, "first_name"),
            "last_name": get(raw, "last_name"),
            "email": email,
            "birth_date": get(raw, "birth_date"),
            "password_hash": pw_hash,
            "created_at": get(raw, "created_at"),
            "last_login": get(raw, "last_login"),
            "failed_attempts": as_int(get(raw, "failed_attempts")),
            "locked_until": get(raw, "locked_until"),
        })

    # every account needs a unique positive ID
    seen, next_id = set(), max([r["id"] for r in out] + [0]) + 1
    for r in out:
        if r["id"] <= 0 or r["id"] in seen:
            r["id"] = next_id
            next_id += 1
        seen.add(r["id"])
    return out


def _save_rows(rows: list):
    """Write the whole list to accounts.xlsx (temp file first, then swap: never half-written)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "users"
    ws.append(ACCOUNT_HEADERS)
    for r in rows:
        ws.append([r.get(f, "") for f in ACCOUNT_FIELDS])

    # Text that starts with "=" must stay text (never run as an Excel formula).
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"

    head_fill = PatternFill("solid", fgColor="2563EB")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [7, 16, 16, 32, 13, 60, 20, 20, 15, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"

    tmp = ACCOUNTS_FILE + ".tmp"
    wb.save(tmp)
    os.replace(tmp, ACCOUNTS_FILE)

    store = _accounts_store()
    store["rows"] = [dict(r) for r in rows]
    store["mtime"] = os.path.getmtime(ACCOUNTS_FILE)


def _import_old_sqlite() -> list:
    """One-time import of the accounts that were kept in perdiapredict.db."""
    if not os.path.exists(DB_FILE):
        return []
    try:
        with closing(sqlite3.connect(DB_FILE, timeout=15)) as conn:
            conn.row_factory = sqlite3.Row
            found = conn.execute("SELECT * FROM users").fetchall()
        return [{
            "id": int(r["id"]),
            "first_name": r["first_name"] or "",
            "last_name": r["last_name"] or "",
            "email": normalize_email(r["email"]),
            "birth_date": r["birth_date"] or "",
            "password_hash": r["password_hash"] or "",
            "created_at": r["created_at"] or "",
            "last_login": r["last_login"] or "",
            "failed_attempts": int(r["failed_attempts"] or 0),
            "locked_until": r["locked_until"] or "",
        } for r in found if r["email"] and r["password_hash"]]
    except Exception:
        return []


def _load_rows() -> list:
    """All accounts (a copy). The Excel file is only re-read when it has changed."""
    store = _accounts_store()
    with store["lock"]:
        if not os.path.exists(ACCOUNTS_FILE):
            old = _import_old_sqlite()
            if old:
                _save_rows(old)
                return [dict(r) for r in old]
            store["rows"], store["mtime"] = [], None
            return []

        mtime = os.path.getmtime(ACCOUNTS_FILE)
        if store["mtime"] != mtime:
            from openpyxl import load_workbook

            wb = load_workbook(ACCOUNTS_FILE, read_only=True, data_only=True)
            try:
                store["rows"] = _rows_from_sheet(wb.active)
            finally:
                wb.close()
            store["mtime"] = mtime
        return [dict(r) for r in store["rows"]]


def _update_account(email: str, **changes):
    store = _accounts_store()
    with store["lock"]:
        rows = _load_rows()
        for r in rows:
            if r["email"] == email:
                r.update(changes)
                _save_rows(rows)
                return


def create_user(first_name: str, last_name: str, email: str, birth_date: date, password: str):
    """Save a new account. Returns (True, None) or (False, translation_key_of_the_error)."""
    email = normalize_email(email)
    try:
        password_hash = hash_password(password)
        store = _accounts_store()
        with store["lock"]:
            rows = _load_rows()
            if any(r["email"] == email for r in rows):
                return False, "auth_email_taken"
            rows.append({
                "id": max([r["id"] for r in rows] + [0]) + 1,
                "first_name": first_name.strip(),
                "last_name": last_name.strip(),
                "email": email,
                "birth_date": birth_date.isoformat(),
                "password_hash": password_hash,
                "created_at": datetime.now().isoformat(timespec="seconds"),
                "last_login": "",
                "failed_attempts": 0,
                "locked_until": "",
            })
            _save_rows(rows)
        return True, None
    except Exception:
        return False, "auth_db_error"


def authenticate(email: str, password: str):
    """Check an email + password against accounts.xlsx.
    Returns (user_dict_or_None, status) where status is "ok", "invalid" or "locked"
    (too many wrong passwords in a row)."""
    email = normalize_email(email)
    now = datetime.now()

    row = next((r for r in _load_rows() if r["email"] == email), None)
    if row is None:
        return None, "invalid"

    if row["locked_until"]:
        try:
            if datetime.fromisoformat(row["locked_until"]) > now:
                return None, "locked"
        except ValueError:
            pass

    if verify_password(password, row["password_hash"]):
        _update_account(
            email, failed_attempts=0, locked_until="",
            last_login=now.isoformat(timespec="seconds"),
        )
        return {
            "id": row["id"],
            "first_name": row["first_name"],
            "last_name": row["last_name"],
            "email": row["email"],
        }, "ok"

    attempts = int(row["failed_attempts"] or 0) + 1
    locked_until, status = "", "invalid"
    if attempts >= MAX_FAILED_LOGINS:
        locked_until = (now + timedelta(minutes=LOCK_MINUTES)).isoformat(timespec="seconds")
        attempts, status = 0, "locked"
    _update_account(email, failed_attempts=attempts, locked_until=locked_until)
    return None, status


def render_accounts_admin():
    """Admin Panel section: the list of accounts, Excel download and Excel restore."""
    rows = _load_rows()
    st.markdown(
        f'<div class="section-title">👥 Registered accounts ({len(rows)})</div>',
        unsafe_allow_html=True,
    )

    if rows:
        table = pd.DataFrame([
            {
                "ID": r["id"], "First name": r["first_name"], "Last name": r["last_name"],
                "Email": r["email"], "Birth date": r["birth_date"],
                "Created at": r["created_at"], "Last login": r["last_login"],
            }
            for r in rows
        ])
        st.dataframe(table, use_container_width=True, hide_index=True)
        with open(ACCOUNTS_FILE, "rb") as f:
            st.download_button(
                "📥 Download accounts (Excel)",
                data=f.read(),
                file_name=ACCOUNTS_FILE,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
                key="dl_accounts",
            )
    else:
        st.info("No accounts yet.")

    with st.expander("♻️ Restore accounts from an Excel backup"):
        st.caption(
            "Upload a previously downloaded accounts.xlsx. Accounts whose email already "
            "exists are skipped, so nothing is overwritten."
        )
        upload = st.file_uploader("accounts.xlsx", type=["xlsx"], key="restore_accounts_file")
        if upload is not None and st.button("Restore now", key="restore_accounts_btn"):
            try:
                from io import BytesIO
                from openpyxl import load_workbook

                wb = load_workbook(BytesIO(upload.getvalue()), read_only=True, data_only=True)
                incoming = _rows_from_sheet(wb.active)
                wb.close()

                store = _accounts_store()
                with store["lock"]:
                    current = _load_rows()
                    known = {r["email"] for r in current}
                    next_id = max([r["id"] for r in current] + [0]) + 1
                    added = skipped = 0
                    for r in incoming:
                        if r["email"] in known or not r["password_hash"].startswith("pbkdf2_sha256$"):
                            skipped += 1
                            continue
                        r["id"] = next_id
                        next_id += 1
                        current.append(r)
                        known.add(r["email"])
                        added += 1
                    if added:
                        _save_rows(current)
                st.success(f"Restored {added} account(s); skipped {skipped}.")
            except Exception as exc:
                st.error(f"Could not read that file: {exc}")
# ACCOUNTS-DB-END'''


def main():
    try:
        text = open(APP_FILE, encoding="utf-8").read()
    except FileNotFoundError:
        sys.exit(f"لم أجد {APP_FILE} في هذا المجلد. ضع هذا الملف بجانبه ثم شغّله.")

    start = text.find("# ACCOUNTS-DB-START")
    end = text.find("# ACCOUNTS-DB-END")
    if start == -1 or end == -1 or end < start:
        sys.exit("لم أجد علامتي ACCOUNTS-DB-START / ACCOUNTS-DB-END في الملف. لم يتم تغيير شيء.")
    end += len("# ACCOUNTS-DB-END")
    text = text[:start] + NEW_BLOCK.strip() + text[end:]

    # Admin Panel: show the accounts section right after the "access granted" message
    if not re.search(r"(?m)^[ \t]*render_accounts_admin\(\)", text):
        text, n = re.subn(
            r'(?m)^([ \t]*)st\.success\(tr\("access"\)\)[ \t]*$',
            lambda m: m.group(0) + "\n" + m.group(1) + "render_accounts_admin()",
            text,
            count=1,
        )
        if n == 0:
            print("تنبيه: لم أجد مكان لوحة الإدارة، فلن يظهر قسم الحسابات هناك (تسجيل الدخول سيعمل بشكل طبيعي).")

    compile(text, APP_FILE, "exec")  # syntax check BEFORE writing anything

    shutil.copyfile(APP_FILE, APP_FILE + ".bak4")
    open(APP_FILE, "w", encoding="utf-8", newline="\n").write(text)
    print("Done. Backup saved as", APP_FILE + ".bak4")


if __name__ == "__main__":
    main()
