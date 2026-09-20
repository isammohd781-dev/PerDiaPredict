"""
apply_update2.py  -  التحديث الثاني

ما الذي يفعله:
  1) زر Enter في أي حقل ينقل المؤشر إلى الحقل التالي، وهكذا حتى آخر حقل
       - في نماذج الدخول والتسجيل: Enter في آخر حقل يرسل النموذج كالمعتاد
       - في نموذج المريض: Enter في آخر حقل ينقل المؤشر إلى زر التنبؤ
       - في القوائم المنسدلة: Enter يختار الخيار الظاهر ثم ينتقل للحقل التالي
       - الترتيب يتبع شكل الصفحة (من الأعلى للأسفل، وفي العربية من اليمين لليسار)
  2) زر "تنزيل جدول النظام الغذائي" (ملف Excel) تحت جدول الأسبوع
  3) حذف كل الإيموجي من الكود ما عدا علامة التحذير وعلامات الصح
       - الأزرار التي كانت إيموجي فقط (الوضع الليلي/النهاري وإعادة التشغيل)
         صارت بأيقونات ستريمليت المدمجة (وليست إيموجي)
       - بدل إيموجي الشعار الاحتياطي (إذا لم يوجد logo.png) صارت الحروف PP

الاستخدام (ضعه بجانب diabetes_app.py ثم شغّله):
    python apply_update2.py

اختياري: لحذف الإيموجي من ملف الترجمات أيضًا:
    python apply_update2.py translations.py

- يحفظ نسخة احتياطية: diabetes_app.py.bak5 (وترجمات: translations.py.bak5)
- يتأكد أن كل موضع تعديل موجود بالعدد المطلوب، وإلا لا يغيّر شيئًا.
- شغّله بعد apply_dob_update.py.
"""
import re
import shutil
import sys

APP_FILE = "diabetes_app.py"

# ---------------------------------------------------------------------------
# 1) Enter = الحقل التالي
# ---------------------------------------------------------------------------
ENTER_NAV = r'''# ENTER-NAV-START
# JavaScript that makes the Enter key jump to the next field. It is added once to
# the page itself (not to the small helper iframe), so it keeps working after reruns.
_ENTER_NAV_JS = r"""
(function () {
  if (window.__ppEnterNav) return;
  window.__ppEnterNav = true;

  var SKIP = ['checkbox', 'radio', 'button', 'submit', 'reset', 'file', 'hidden', 'image', 'range', 'color'];
  var BLOCK = '[data-testid="stElementContainer"], [data-testid="element-container"], .element-container';

  function isField(inp) {
    if (!inp || inp.tagName !== 'INPUT') return false;
    var type = (inp.getAttribute('type') || 'text').toLowerCase();
    if (SKIP.indexOf(type) !== -1) return false;
    if (inp.disabled) return false;
    return inp.offsetParent !== null;
  }

  // All fields of a card / form in the order the eye reads them:
  // top to bottom, and inside one row left to right (right to left in Arabic).
  function fieldsIn(scope) {
    var list = Array.prototype.filter.call(scope.querySelectorAll('input'), isField);
    var rtl = window.getComputedStyle(scope).direction === 'rtl';
    var items = list.map(function (inp) {
      var box = inp.closest(BLOCK) || inp;
      var r = box.getBoundingClientRect();
      return { el: inp, top: r.top, left: r.left };
    });
    items.sort(function (a, b) { return a.top - b.top; });
    var rows = [], cur = null;
    items.forEach(function (it) {
      if (cur && Math.abs(it.top - cur.top) < 32) { cur.items.push(it); }
      else { cur = { top: it.top, items: [it] }; rows.push(cur); }
    });
    var out = [];
    rows.forEach(function (row) {
      row.items.sort(function (a, b) { return rtl ? b.left - a.left : a.left - b.left; });
      row.items.forEach(function (it) { out.push(it.el); });
    });
    return out;
  }

  document.addEventListener('keydown', function (e) {
    if (e.key !== 'Enter' || e.shiftKey || e.ctrlKey || e.altKey || e.metaKey || e.isComposing) return;
    var el = e.target;
    if (!isField(el)) return;

    var scope = el.closest('[data-testid="stForm"]') ||
                el.closest('.st-key-patient_card') ||
                el.closest('[data-testid="stMainBlockContainer"]');
    if (!scope) return;

    var inForm = scope.getAttribute('data-testid') === 'stForm';
    var predictBtn = function () { return scope.querySelector('.st-key-predict_btn button'); };
    var fields = fieldsIn(scope);
    var idx = fields.indexOf(el);
    if (idx === -1) return;

    var isLast = idx === fields.length - 1;
    // Last field of a form: normal Enter (submits the form).
    // Last field elsewhere with nothing to jump to: normal Enter.
    if (isLast && (inForm || !predictBtn())) return;

    function go() {
      var next = fieldsIn(scope)[idx + 1];
      if (next) {
        next.focus();
        try { next.select(); } catch (_) {}
      } else {
        var b = predictBtn();
        if (b) b.focus();
      }
    }

    if (el.closest('[data-baseweb="select"]')) {
      // Drop-down field: let Enter pick the highlighted option first, then move on.
      setTimeout(go, 90);
    } else {
      e.preventDefault();
      e.stopPropagation();   // stops Streamlit from submitting the form on this Enter
      go();
    }
  }, true);
})();
"""


def inject_enter_navigation():
    """Adds the Enter-key navigation script to the page (once)."""
    import json

    components.html(
        "<script>(function(){"
        "var d=window.parent.document;"
        "if(d.getElementById('pp-enter-nav'))return;"
        "var s=d.createElement('script');"
        "s.id='pp-enter-nav';"
        "s.text=" + json.dumps(_ENTER_NAV_JS) + ";"
        "d.head.appendChild(s);"
        "})();</script>",
        height=0,
    )
# ENTER-NAV-END


'''

# ---------------------------------------------------------------------------
# 2) تنزيل جدول النظام الغذائي
# ---------------------------------------------------------------------------
MEAL_TEXT = r'''for _code, _vals in _AUTH_UI_TEXT.items():
    EXTRA_TEXT.setdefault(_code, {}).update(_vals)

# Label of the meal-plan download button (other languages fall back to English).
_MEAL_DL_TEXT = {
    "en": {"download_meal_plan": "Download the meal plan (Excel)"},
    "ar": {"download_meal_plan": "تنزيل جدول النظام الغذائي (Excel)"},
    "es": {"download_meal_plan": "Descargar el plan de comidas (Excel)"},
}
for _code, _vals in _MEAL_DL_TEXT.items():
    EXTRA_TEXT.setdefault(_code, {}).update(_vals)
'''

MEAL_FUNC = r'''def _meal_plan_excel(df: pd.DataFrame, rtl: bool = False) -> bytes:
    """The weekly meal plan as an .xlsx file (bytes)."""
    import io
    from openpyxl.styles import Alignment, Font, PatternFill

    sheet = re.sub(r"[\\/*?:\[\]]", " ", str(tr("meal_plan"))).strip()[:31] or "Meal plan"
    buffer = io.BytesIO()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name=sheet)
        ws = writer.sheets[sheet]
        if rtl:
            ws.sheet_view.rightToLeft = True

        head_fill = PatternFill("solid", fgColor="2563EB")
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 24

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top", wrap_text=True, horizontal="right" if rtl else "left"
                )
            row[0].font = Font(bold=True)

        ws.column_dimensions["A"].width = 12
        for letter in "BCDE":
            ws.column_dimensions[letter].width = 38
    return buffer.getvalue()


def render_meal_plan(veg: bool = False):'''

MEAL_BUTTON_OLD = "    st.dataframe(df, use_container_width=True, hide_index=True)\n"
MEAL_BUTTON_NEW = '''    st.dataframe(df, use_container_width=True, hide_index=True)

    st.download_button(
        tr("download_meal_plan"),
        data=_meal_plan_excel(df, is_rtl()),
        file_name="Weekly_Meal_Plan_" + ("veg" if veg else "nonveg") + ".xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="meal_plan_download",
        use_container_width=True,
    )
'''

# (النص القديم، النص الجديد) - كل نص قديم يجب أن يظهر مرة واحدة بالضبط
REPLACEMENTS = [
    # Enter navigation: الدالة قبل render_footer، وتُستدعى في بدايتها (كل الصفحات ما عدا الشعار)
    ("def render_footer():", ENTER_NAV + "def render_footer():\n    inject_enter_navigation()"),
    # نص زر التنزيل
    (
        "for _code, _vals in _AUTH_UI_TEXT.items():\n    EXTRA_TEXT.setdefault(_code, {}).update(_vals)\n",
        MEAL_TEXT,
    ),
    # دالة Excel
    ("def render_meal_plan(veg: bool = False):", MEAL_FUNC),
    # زر التنزيل تحت الجدول
    (MEAL_BUTTON_OLD, MEAL_BUTTON_NEW),
]

# ---------------------------------------------------------------------------
# 3) حذف الإيموجي
# ---------------------------------------------------------------------------
# حالات خاصة أولًا: (نمط، بديل، العدد المطلوب)
SPECIAL = [
    # أيقونة تبويب المتصفح الاحتياطية
    (r'(_page_icon = "logo\.png" if os\.path\.exists\("logo\.png"\) else )"\U0001FA7A\uFE0F?"',
     r"\1None", 1),
    # الشعار الاحتياطي في شاشة البداية والرأس
    (r'\b(logo_html|brand_icon_html) = "\U0001FA7A\uFE0F?"', r'\1 = "PP"', 2),
    # الشعار الاحتياطي في صفحات الدخول
    (r'(<span class="auth-logo-fallback">)\U0001FA7A\uFE0F?(</span>)', r"\1PP\2", 1),
    # زر إعادة التشغيل (كان إيموجي فقط)
    (r'"\U0001F504\uFE0F?",\n([ \t]+)key="menu_restart",',
     '"",\n\\1icon=":material/refresh:",\n\\1key="menu_restart",', 1),
    # زر الوضع الليلي/النهاري (كان إيموجي فقط)
    (r'"\u2600\uFE0F?" if is_dark else "\U0001F319\uFE0F?",\n([ \t]+)key="theme_btn",',
     '"",\n\\1icon=":material/light_mode:" if is_dark else ":material/dark_mode:",\n\\1key="theme_btn",', 1),
    # أيقونات عناوين الأسئلة الخاصة بالجنس/الأطفال
    (r"(?m)^[ \t]*gender_icon = [^\n]*\n", "", 3),
    (r"\{gender_icon\} \{gender_title\}", "{gender_title}", 1),
]

# ما يُحذف: كل الإيموجي، ما عدا  ⚠ (تحذير)  و ✅ ✔ ✓ (صح)
_EMOJI_RE = re.compile(
    r"([\u26A0\u2705\u2714\u2713]\uFE0F?)"                                   # يُبقى كما هو
    r"|([\U0001F000-\U0001FAFF\u2600-\u27BF\u2B00-\u2BFF\u25B6\u25C0]\uFE0F?[ ]?)"  # يُحذف
)


def strip_emoji(text: str):
    removed = 0

    def _sub(match):
        nonlocal removed
        if match.group(1):
            return match.group(1)
        removed += 1
        return ""

    return _EMOJI_RE.sub(_sub, text), removed


def main():
    extra_files = [a for a in sys.argv[1:] if a.endswith(".py")]

    try:
        text = open(APP_FILE, encoding="utf-8").read()
    except FileNotFoundError:
        sys.exit(f"لم أجد {APP_FILE} في هذا المجلد. ضع هذا الملف بجانبه ثم شغّله.")

    if "# ENTER-NAV-START" in text:
        sys.exit("هذا التحديث مطبّق مسبقًا على الملف. لم يتم تغيير شيء.")

    # ---- تحقق من كل المواضع قبل أي تغيير ----
    for old, _new in REPLACEMENTS:
        count = text.count(old)
        if count != 1:
            sys.exit(f"الموضع التالي وُجد {count} مرة (المطلوب مرة واحدة). لم يتم تغيير شيء:\n{old[:120]}")
    for pattern, _repl, expected in SPECIAL:
        found = len(re.findall(pattern, text))
        if found != expected:
            sys.exit(f"النمط التالي وُجد {found} مرة (المطلوب {expected}). لم يتم تغيير شيء:\n{pattern[:120]}")

    # ---- التطبيق ----
    for old, new in REPLACEMENTS:
        text = text.replace(old, new, 1)
    for pattern, repl, _expected in SPECIAL:
        text = re.sub(pattern, repl, text)
    text, removed = strip_emoji(text)

    compile(text, APP_FILE, "exec")  # syntax check BEFORE writing anything

    # ملفات إضافية (الترجمات): حذف الإيموجي فقط
    extra_results = []
    for name in extra_files:
        try:
            src = open(name, encoding="utf-8").read()
        except FileNotFoundError:
            print(f"تخطّيت {name}: الملف غير موجود.")
            continue
        new_src, n = strip_emoji(src)
        compile(new_src, name, "exec")
        extra_results.append((name, src, new_src, n))

    shutil.copyfile(APP_FILE, APP_FILE + ".bak5")
    open(APP_FILE, "w", encoding="utf-8", newline="\n").write(text)
    print(f"Done. {APP_FILE}: removed {removed} emoji. Backup: {APP_FILE}.bak5")

    for name, src, new_src, n in extra_results:
        shutil.copyfile(name, name + ".bak5")
        open(name, "w", encoding="utf-8", newline="\n").write(new_src)
        print(f"Done. {name}: removed {n} emoji. Backup: {name}.bak5")


if __name__ == "__main__":
    main()